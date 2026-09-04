"""
test_pipeline.py

End-to-end test harness for the PDF -> chunks -> AI plan -> PowerPoint pipeline.

This script tests each component independently first, then runs the complete
pipeline if the earlier stages succeed.

Expected project structure:

your_project/
├── pdf_extractor.py
├── chunking.py
├── presentation_ai.py
├── ppt_generator.py
├── test_pipeline.py
├── input/
│   └── osi_model.pdf
├── templates/
│   └── template.pptx
└── output/

Usage:
    python test_pipeline.py

Or specify your own files:
    python test_pipeline.py input/my.pdf templates/my_template.pptx

Notes:
- The script intentionally imports your existing modules rather than duplicating
  their logic.
- Because function signatures can vary between versions of the modules, the
  test runner tries several common calling conventions and reports failures
  clearly instead of silently guessing.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import traceback
from pathlib import Path
from typing import Any, Callable
print("test_pipeline.py started", flush=True)


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

DEFAULT_PDF = Path("input/osi_model.pdf")
DEFAULT_TEMPLATE = Path("templates/template.pptx")
OUTPUT_DIR = Path("output")

EXTRACTION_JSON = OUTPUT_DIR / "test_extraction.json"
CHUNKS_JSON = OUTPUT_DIR / "test_chunks.json"
PLAN_JSON = OUTPUT_DIR / "presentation_plan.json"
PPTX_OUTPUT = OUTPUT_DIR / "test_presentation.pptx"
IMAGE_DIR = OUTPUT_DIR / "osi_images"


# ---------------------------------------------------------------------------
# Console helpers
# ---------------------------------------------------------------------------

def print_header(title: str) -> None:
    print("\n" + "=" * 78)
    print(title)
    print("=" * 78)


def print_ok(message: str) -> None:
    print(f"[PASS] {message}")


def print_fail(message: str) -> None:
    print(f"[FAIL] {message}")


def print_warn(message: str) -> None:
    print(f"[WARN] {message}")


def print_info(message: str) -> None:
    print(f"[INFO] {message}")


# ---------------------------------------------------------------------------
# General helpers
# ---------------------------------------------------------------------------

def save_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False, default=str)


def load_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def import_module_safely(module_name: str):
    try:
        return __import__(module_name)
    except Exception as exc:
        print_fail(f"Could not import {module_name}: {exc}")
        traceback.print_exc()
        return None


def call_variants(
    function: Callable[..., Any],
    variants: list[tuple[str, tuple[Any, ...], dict[str, Any]]],
) -> tuple[Any | None, str | None]:
    """
    Try several likely function signatures.

    Returns:
        (result, description) on success
        (None, None) if every variant failed
    """
    last_error = None

    for description, args, kwargs in variants:
        try:
            result = function(*args, **kwargs)
            print_info(f"Called {function.__name__} using: {description}")
            return result, description
        except TypeError as exc:
            # TypeError is commonly caused by a mismatched signature.
            last_error = exc
            continue
        except Exception as exc:
            # The signature was accepted, but the function itself failed.
            print_fail(
                f"{function.__name__} failed while using {description}: {exc}"
            )
            traceback.print_exc()
            return None, None

    if last_error:
        print_fail(
            f"No compatible calling convention found for "
            f"{function.__name__}: {last_error}"
        )

    return None, None


def first_existing_path(*paths: Path) -> Path | None:
    for path in paths:
        if path.exists():
            return path
    return None


# ---------------------------------------------------------------------------
# Test 1: PDF extractor
# ---------------------------------------------------------------------------

def test_pdf_extractor(pdf_path: Path) -> dict[str, Any] | None:
    print_header("TEST 1 — PDF EXTRACTION")

    module = import_module_safely("pdf_extractor")
    if module is None:
        return None

    if not hasattr(module, "extract_pdf"):
        print_fail("pdf_extractor.py does not contain extract_pdf().")
        return None

    extract_pdf = module.extract_pdf

    variants = [
        (
            "extract_pdf(pdf_path, image_output_dir=..., extract_pdf_images=True)",
            (str(pdf_path),),
            {
                "image_output_dir": str(IMAGE_DIR),
                "extract_pdf_images": True,
            },
        ),
        (
            "extract_pdf(pdf_path, image_output_dir=...)",
            (str(pdf_path),),
            {
                "image_output_dir": str(IMAGE_DIR),
            },
        ),
        (
            "extract_pdf(pdf_path)",
            (str(pdf_path),),
            {},
        ),
        (
            "extract_pdf(pdf_path=..., image_output_dir=..., extract_pdf_images=True)",
            (),
            {
                "pdf_path": str(pdf_path),
                "image_output_dir": str(IMAGE_DIR),
                "extract_pdf_images": True,
            },
        ),
    ]

    result, _ = call_variants(extract_pdf, variants)

    if result is None:
        return None

    if not isinstance(result, dict):
        print_fail(
            f"extract_pdf() returned {type(result).__name__}, expected dict."
        )
        return None

    save_json(EXTRACTION_JSON, result)

    pages = result.get("pages", [])
    paragraphs = result.get("paragraphs", [])
    images = result.get("images", [])
    images_by_page = result.get("images_by_page", {})

    print_ok(f"Extraction returned a dictionary.")
    print_info(f"Pages: {len(pages) if isinstance(pages, list) else 'unknown'}")
    print_info(
        f"Paragraphs: "
        f"{len(paragraphs) if isinstance(paragraphs, list) else 'unknown'}"
    )
    print_info(
        f"Images: {len(images) if isinstance(images, list) else 'unknown'}"
    )
    print_info(
        f"Images by page entries: "
        f"{len(images_by_page) if isinstance(images_by_page, dict) else 'unknown'}"
    )
    print_ok(f"Saved extraction JSON: {EXTRACTION_JSON}")

    # Check the image metadata expected by the revised extractor.
    if isinstance(images, list) and images:
        sample = images[0]
        expected_keys = {
            "image_id",
            "page",
            "path",
            "width",
            "height",
            "description",
        }
        missing = expected_keys - set(sample.keys())

        if missing:
            print_warn(
                "Image metadata is missing expected fields: "
                + ", ".join(sorted(missing))
            )
        else:
            print_ok(
                "Image metadata contains image_id, page, path, dimensions, "
                "and description."
            )

        described = sum(
            1
            for image in images
            if isinstance(image, dict) and image.get("description")
        )
        print_info(f"Images with descriptions: {described}/{len(images)}")

    else:
        print_warn("No PDF images were extracted.")

    return result


# ---------------------------------------------------------------------------
# Test 2: Chunking
# ---------------------------------------------------------------------------

def test_chunking(
    pdf_path: Path,
    extraction: dict[str, Any],
) -> Any | None:
    print_header("TEST 2 — CHUNKING")

    module = import_module_safely("chunking")
    if module is None:
        return None

    # Prefer the function mentioned in the pipeline design.
    candidate_names = [
        "pdf_to_embedded_chunks",
        "create_chunks",
        "chunk_pdf",
        "build_chunks",
        "make_chunks",
    ]

    function_name = next(
        (name for name in candidate_names if hasattr(module, name)),
        None,
    )

    if function_name is None:
        print_fail(
            "Could not find a known chunking function in chunking.py. "
            f"Tried: {', '.join(candidate_names)}"
        )
        return None

    function = getattr(module, function_name)
    print_info(f"Using chunking function: {function_name}()")

    variants = [
        (
            f"{function_name}(extraction)",
            (extraction,),
            {},
        ),
        (
            f"{function_name}(pdf_path)",
            (str(pdf_path),),
            {},
        ),
        (
            f"{function_name}(pdf_path, extraction)",
            (str(pdf_path), extraction),
            {},
        ),
        (
            f"{function_name}(extraction=extraction)",
            (),
            {"extraction": extraction},
        ),
        (
            f"{function_name}(pdf_path=..., extraction=extraction)",
            (),
            {
                "pdf_path": str(pdf_path),
                "extraction": extraction,
            },
        ),
    ]

    result, _ = call_variants(function, variants)

    if result is None:
        return None

    # Some chunking implementations return a wrapper dictionary.
    if isinstance(result, dict):
        chunks = (
            result.get("chunks")
            or result.get("embedded_chunks")
            or result.get("data")
            or result
        )
    else:
        chunks = result

    if not isinstance(chunks, (list, tuple)):
        print_fail(
            f"Chunking returned {type(result).__name__}; "
            "could not identify a list of chunks."
        )
        return None

    chunks = list(chunks)
    save_json(CHUNKS_JSON, result)

    print_ok(f"Chunking returned {len(chunks)} chunks.")
    print_ok(f"Saved chunks JSON: {CHUNKS_JSON}")

    if chunks:
        print_info("First chunk preview:")
        preview = chunks[0]

        if isinstance(preview, dict):
            print(json.dumps(preview, indent=2, ensure_ascii=False)[:1200])
        else:
            print(str(preview)[:1200])

    return result


# ---------------------------------------------------------------------------
# Test 3: Presentation AI
# ---------------------------------------------------------------------------

def test_presentation_ai(
    extraction: dict[str, Any],
    chunks: Any,
) -> dict[str, Any] | None:
    print_header("TEST 3 — AI PRESENTATION PLANNING")

    module = import_module_safely("presentation_ai")
    if module is None:
        return None

    candidate_names = [
        "create_presentation_plan",
        "generate_presentation_plan",
        "build_presentation_plan",
        "plan_presentation",
        "make_presentation_plan",
    ]

    function_name = next(
        (name for name in candidate_names if hasattr(module, name)),
        None,
    )

    if function_name is None:
        print_fail(
            "Could not find a known presentation-planning function in "
            "presentation_ai.py. "
            f"Tried: {', '.join(candidate_names)}"
        )
        return None

    function = getattr(module, function_name)
    print_info(f"Using AI planning function: {function_name}()")

    variants = [
        (
            f"{function_name}(chunks, extraction)",
            (chunks, extraction),
            {},
        ),
        (
            f"{function_name}(chunks)",
            (chunks,),
            {},
        ),
        (
            f"{function_name}(extraction, chunks)",
            (extraction, chunks),
            {},
        ),
        (
            f"{function_name}(chunks=chunks, extraction=extraction)",
            (),
            {
                "chunks": chunks,
                "extraction": extraction,
            },
        ),
        (
            f"{function_name}(chunks=chunks)",
            (),
            {"chunks": chunks},
        ),
    ]

    result, _ = call_variants(function, variants)

    if result is None:
        return None

    # The AI function may return a JSON string rather than a Python dict.
    if isinstance(result, str):
        try:
            result = json.loads(result)
        except json.JSONDecodeError:
            print_fail(
                "Presentation AI returned a string that is not valid JSON."
            )
            print_info(result[:2000])
            return None

    if not isinstance(result, dict):
        print_fail(
            f"Presentation AI returned {type(result).__name__}, expected dict."
        )
        return None

    save_json(PLAN_JSON, result)

    slides = result.get("slides", [])

    print_ok("Presentation plan returned as a dictionary.")
    print_info(
        f"Planned slides: "
        f"{len(slides) if isinstance(slides, list) else 'unknown'}"
    )
    print_ok(f"Saved presentation plan: {PLAN_JSON}")

    # Check common plan fields.
    expected_plan_fields = ["slides"]
    missing = [
        field for field in expected_plan_fields if field not in result
    ]

    if missing:
        print_warn(
            "Plan is missing expected fields: " + ", ".join(missing)
        )

    if isinstance(slides, list) and slides:
        print_info("First slide plan:")
        print(json.dumps(slides[0], indent=2, ensure_ascii=False)[:2000])

        visual_slides = 0
        visual_refs = 0

        for slide in slides:
            if not isinstance(slide, dict):
                continue

            slide_type = str(
                slide.get("type", slide.get("slide_type", ""))
            ).lower()

            if slide_type in {"visual", "image", "two_column", "content_visual"}:
                visual_slides += 1

            for key in ("image_id", "visual_id", "images", "visuals"):
                value = slide.get(key)
                if value:
                    if isinstance(value, list):
                        visual_refs += len(value)
                    else:
                        visual_refs += 1

        print_info(f"Slides containing visual-oriented content: {visual_slides}")
        print_info(f"Visual references found in plan: {visual_refs}")

        if len(slides) < 15:
            print_warn(
                "The generated plan contains fewer than 15 slides. "
                "If your target is ~20 slides, inspect presentation_ai.py."
            )
        elif len(slides) > 25:
            print_warn(
                "The generated plan contains more than 25 slides. "
                "If your target is ~20 slides, inspect presentation_ai.py."
            )
        else:
            print_ok("Slide count is within the expected ~20-slide range.")

    return result


# ---------------------------------------------------------------------------
# Test 4: PowerPoint generator
# ---------------------------------------------------------------------------

def test_ppt_generator(
    plan_path: Path,
    template_path: Path | None,
) -> bool:
    print_header("TEST 4 — POWERPOINT GENERATION")

    module = import_module_safely("ppt_generator")
    if module is None:
        return False

    # The generic generator created earlier is normally runnable as a CLI.
    # Running it as a subprocess avoids making assumptions about its internal
    # function names/signature.
    if not hasattr(module, "main"):
        print_warn(
            "ppt_generator.py does not expose main(); trying CLI execution."
        )

    command = [
        sys.executable,
        "ppt_generator.py",
        str(plan_path),
        str(PPTX_OUTPUT),
    ]

    if template_path is not None:
        command.extend(["--template", str(template_path)])

    # Image root is useful if extracted images are referenced by relative paths.
    command.extend(["--image-root", str(IMAGE_DIR)])

    print_info("Running:")
    print_info(" ".join(command))

    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            check=False,
        )
    except Exception as exc:
        print_fail(f"Could not launch ppt_generator.py: {exc}")
        return False

    if completed.stdout:
        print("\n--- ppt_generator.py output ---")
        print(completed.stdout)

    if completed.stderr:
        print("\n--- ppt_generator.py errors/warnings ---")
        print(completed.stderr)

    if completed.returncode != 0:
        print_fail(
            f"ppt_generator.py exited with code {completed.returncode}."
        )
        return False

    if not PPTX_OUTPUT.exists():
        print_fail(
            f"ppt_generator.py completed successfully, but the expected "
            f"PowerPoint was not created: {PPTX_OUTPUT}"
        )
        return False

    print_ok(f"PowerPoint created: {PPTX_OUTPUT}")
    print_info(f"File size: {PPTX_OUTPUT.stat().st_size:,} bytes")

    # Optional validation with python-pptx.
    try:
        from pptx import Presentation

        prs = Presentation(str(PPTX_OUTPUT))
        print_info(f"Generated PowerPoint contains {len(prs.slides)} slides.")

        if len(prs.slides) == 0:
            print_fail("Generated PowerPoint contains zero slides.")
            return False

        print_ok("PowerPoint can be opened successfully by python-pptx.")
    except ImportError:
        print_warn(
            "python-pptx is not installed, so the generated PPTX could not "
            "be opened for validation."
        )
    except Exception as exc:
        print_fail(f"Generated PPTX could not be validated: {exc}")
        return False

    return True


# ---------------------------------------------------------------------------
# Dependency check
# ---------------------------------------------------------------------------

def check_dependencies() -> bool:
    print_header("DEPENDENCY CHECK")

    packages = [
        ("fitz", "PyMuPDF"),
        ("pptx", "python-pptx"),
        ("PIL", "Pillow"),
    ]

    all_ok = True

    for import_name, package_name in packages:
        try:
            __import__(import_name)
            print_ok(package_name)
        except ImportError:
            print_fail(
                f"{package_name} is not installed. "
                f"Install it with: pip install {package_name}"
            )
            all_ok = False

    return all_ok


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> int:
    print_header("PDF → PRESENTATION PIPELINE TEST")

    # Optional positional arguments.
    pdf_path = Path(sys.argv[1]) if len(sys.argv) >= 2 else DEFAULT_PDF
    template_path = (
        Path(sys.argv[2]) if len(sys.argv) >= 3 else DEFAULT_TEMPLATE
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)

    print_info(f"PDF: {pdf_path}")
    print_info(f"Template: {template_path}")
    print_info(f"Output directory: {OUTPUT_DIR}")

    if not pdf_path.exists():
        print_fail(f"PDF not found: {pdf_path}")
        print_info(
            "Put your PDF at input/osi_model.pdf or run:\n"
            "python test_pipeline.py path/to/your.pdf"
        )
        return 1

    if not template_path.exists():
        print_warn(
            f"PowerPoint template not found: {template_path}\n"
            "The extraction, chunking, and AI tests can still run. "
            "The PowerPoint generation test will be skipped."
        )
        template_path_for_test = None
    else:
        template_path_for_test = template_path

    if not check_dependencies():
        print_fail("Required Python dependencies are missing.")
        return 1

    results: dict[str, bool] = {
        "pdf_extractor": False,
        "chunking": False,
        "presentation_ai": False,
        "ppt_generator": False,
    }

    # ---------------------------------------------------------------
    # 1. PDF extraction
    # ---------------------------------------------------------------

    extraction = test_pdf_extractor(pdf_path)

    if extraction is None:
        print_fail("Stopping because PDF extraction failed.")
        print_summary(results)
        return 1

    results["pdf_extractor"] = True

    # ---------------------------------------------------------------
    # 2. Chunking
    # ---------------------------------------------------------------

    chunks = test_chunking(pdf_path, extraction)

    if chunks is None:
        print_fail("Stopping because chunking failed.")
        print_summary(results)
        return 1

    results["chunking"] = True

    # ---------------------------------------------------------------
    # 3. AI presentation planning
    # ---------------------------------------------------------------

    plan = test_presentation_ai(extraction, chunks)

    if plan is None:
        print_fail(
            "Stopping because presentation AI planning failed.\n"
            "This stage may require your AI API credentials/configuration."
        )
        print_summary(results)
        return 1

    results["presentation_ai"] = True

    # ---------------------------------------------------------------
    # 4. PowerPoint generation
    # ---------------------------------------------------------------

    if template_path_for_test is None:
        print_warn("Skipping PowerPoint generation because no template exists.")
    else:
        ppt_ok = test_ppt_generator(PLAN_JSON, template_path_for_test)
        results["ppt_generator"] = ppt_ok

    # ---------------------------------------------------------------
    # Final summary
    # ---------------------------------------------------------------

    print_summary(results)

    if all(results.values()):
        print_ok("ALL PIPELINE COMPONENTS PASSED.")
        return 0

    if results["pdf_extractor"] and results["chunking"] and results["presentation_ai"]:
        print_warn(
            "Extraction, chunking, and AI planning passed. "
            "PowerPoint generation did not pass."
        )
        return 1

    return 1


def print_summary(results: dict[str, bool]) -> None:
    print_header("TEST SUMMARY")

    labels = {
        "pdf_extractor": "1. PDF extractor",
        "chunking": "2. Chunking",
        "presentation_ai": "3. Presentation AI",
        "ppt_generator": "4. PowerPoint generator",
    }

    for key, label in labels.items():
        status = "PASS" if results[key] else "FAIL"
        print(f"{status:>5}  {label}")

    print("\nGenerated test artifacts:")
    for path in [
        EXTRACTION_JSON,
        CHUNKS_JSON,
        PLAN_JSON,
        PPTX_OUTPUT,
    ]:
        if path.exists():
            print(f"  - {path}")


if __name__ == "__main__":
    # raise SystemExit(main())
    main()
